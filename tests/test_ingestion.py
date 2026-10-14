import pytest
import io
from pathlib import Path
from app.document_processor import process_document, compute_sha256, chunk_text_blocks

def test_txt_ingestion():
    content = (
        "## Academic Calendar\n"
        "The Fall semester starts on August 28. Students may add courses during the first two weeks.\n\n"
        "## Final Exams\n"
        "Final examinations are scheduled from December 11 through December 16 in assigned lecture halls."
    ).encode("utf-8")

    chunks, file_type = process_document("calendar.txt", content)
    assert file_type == "TXT"
    assert len(chunks) >= 2
    assert chunks[0]["page_number"] == 1
    assert "Academic Calendar" in chunks[0]["section_heading"]
    assert "August 28" in chunks[0]["content"]

def test_empty_file_rejection():
    with pytest.raises(ValueError, match="empty"):
        process_document("empty.txt", b"")

def test_unsupported_format_rejection():
    with pytest.raises(ValueError, match="Unsupported file format"):
        process_document("program.exe", b"binary content")

def test_sha256_computation():
    content = b"Sample university policy text."
    h1 = compute_sha256(content)
    h2 = compute_sha256(content)
    assert h1 == h2
    assert len(h1) == 64

def test_pdf_ingestion(tmp_path):
    import fitz
    doc = fitz.open()
    p = doc.new_page()
    rect = fitz.Rect(50, 50, 500, 500)
    p.insert_textbox(rect, "SECTION: Library Access\nThe central library is open 24 hours during finals week with valid student badge.", fontsize=12)
    pdf_bytes = doc.tobytes()
    doc.close()

    chunks, file_type = process_document("library.pdf", pdf_bytes)
    assert file_type == "PDF"
    assert len(chunks) >= 1
    assert chunks[0]["page_number"] == 1
    assert "central library" in chunks[0]["content"]

def test_docx_ingestion():
    import docx
    doc = docx.Document()
    doc.add_heading("Housing Policies", level=1)
    doc.add_paragraph("Quiet hours are observed from 10 PM to 8 AM Sunday through Thursday in all residential dormitories.")
    bio = io.BytesIO()
    doc.save(bio)
    docx_bytes = bio.getvalue()

    chunks, file_type = process_document("housing.docx", docx_bytes)
    assert file_type == "DOCX"
    assert len(chunks) >= 1
    assert chunks[0]["section_heading"] == "Housing Policies"
    assert "Quiet hours" in chunks[0]["content"]
