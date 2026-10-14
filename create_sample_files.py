import docx
from pathlib import Path
import fitz  # PyMuPDF

SAMPLE_DIR = Path(__file__).resolve().parent / "sample_docs"
SAMPLE_DIR.mkdir(parents=True, exist_ok=True)

def generate_docx():
    doc = docx.Document()
    doc.add_heading("SAMPLE UNIVERSITY — RESIDENCE LIFE & CAMPUS DINING", level=0)
    doc.add_paragraph("[Notice: Fictional campus document for demonstration and testing.]")

    doc.add_heading("Section 1: Move-In and Move-Out Guidelines", level=1)
    doc.add_paragraph(
        "Fall semester move-in occurs during the weekend prior to the start of classes. First-year students "
        "move in on Thursday, August 24, according to assigned staggered arrival time slots (8:00 AM to 4:00 PM). "
        "Returning upperclass students move in on Saturday, August 26. All students must complete online check-in "
        "via the Housing Portal 48 hours prior to arrival."
    )
    doc.add_paragraph(
        "Residence halls close at 12:00 PM on the Saturday following final examinations. Students needing late departure "
        "must submit a Late Stay Request form at least one week in advance. Unauthorized late departures incur a $75 fee."
    )

    doc.add_heading("Section 2: Room Change Requests and Quiet Hours", level=1)
    doc.add_paragraph(
        "A two-week Room Freeze is observed at the beginning of each semester during which no room changes are permitted. "
        "Starting Monday of the third week of classes, the Room Change Request portal opens. Students must first meet with "
        "their Resident Assistant (RA) to participate in a mediated roommate agreement before submitting a transfer request."
    )
    doc.add_paragraph(
        "Quiet hours are strictly enforced in all residential communities: Sunday through Thursday from 10:00 PM to 8:00 AM, "
        "and Friday through Saturday from 12:00 AM (midnight) to 9:00 AM. Courtesy hours are in effect 24 hours a day."
    )

    doc.add_heading("Section 3: Meal Plan Options and Dining Dollar Rollover", level=1)
    doc.add_paragraph(
        "All first-year residential students are required to select either the Unlimited Access or 14-Meal Plan. "
        "Sophomores, juniors, and commuter students may choose any plan or a block plan. Students wishing to downgrade "
        "or upgrade their meal plan must submit changes before the second Friday of the semester (September 8 at 5:00 PM). "
        "After this cutoff date, no meal plan cancellations or downgrades are accepted."
    )
    doc.add_paragraph(
        "Dining Dollars roll over from the Fall to Spring semester provided the student maintains an active meal plan for "
        "both terms. All unused Dining Dollars expire on the final day of Spring semester final exams and are non-refundable."
    )

    # Add a table
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Light Shading'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Meal Plan'
    hdr_cells[1].text = 'Weekly Meals'
    hdr_cells[2].text = 'Dining Dollars / Sem'

    data = [
        ('Unlimited Access', 'Unlimited swipes', '$250 Dining Dollars'),
        ('14-Meal Plan', '14 meals / week', '$400 Dining Dollars'),
        ('Block 75 (Commuter)', '75 meals total', '$500 Dining Dollars')
    ]
    for plan, meals, dd in data:
        row_cells = table.add_row().cells
        row_cells[0].text = plan
        row_cells[1].text = meals
        row_cells[2].text = dd

    docx_path = SAMPLE_DIR / "housing_and_dining_policies.docx"
    doc.save(str(docx_path))
    print(f"Generated DOCX: {docx_path}")

def generate_pdf():
    doc = fitz.open()

    # Page 1: FAFSA, Payment Due Dates, Payment Plans
    p1 = doc.new_page(width=595, height=842)  # A4
    text_p1 = (
        "SAMPLE UNIVERSITY — FINANCIAL AID & STUDENT ACCOUNTS\n"
        "Frequently Asked Questions (Tuition, Financial Aid & Scholarships)\n"
        "[Notice: Fictional campus document for demonstration and testing.]\n\n"
        "SECTION: FAFSA Application and Priority Deadlines\n"
        "The Free Application for Federal Student Aid (FAFSA) opens on October 1 each year for the subsequent "
        "academic year. Sample University's institutional priority filing deadline is February 15. Submitting "
        "the FAFSA by this date ensures maximum consideration for need-based institutional grants and Federal "
        "Work-Study funds. Students submitting after February 15 are considered on a rolling, funds-available basis.\n\n"
        "SECTION: Tuition Payment Deadlines and Late Fees\n"
        "Fall semester tuition and mandatory fees must be paid in full or covered by an approved payment plan "
        "by August 15. For the Spring semester, the billing deadline is January 5. A late payment fee of $100 "
        "is assessed if payment is not settled within 5 business days after the due date. Accounts with unpaid "
        "balances exceeding $500 will have a financial hold placed, preventing registration and transcript release.\n\n"
        "SECTION: Monthly Tuition Installment Plans\n"
        "The Bursar's Office provides interest-free 4-month and 5-month payment plans for each semester. "
        "An enrollment fee of $45 per semester applies. Payments are automatically debited on the 1st of each "
        "month via ACH bank transfer or credit card (a 2.85% convenience fee applies to credit card payments)."
    )
    rect = fitz.Rect(50, 50, 545, 780)
    p1.insert_textbox(rect, text_p1, fontsize=11, fontname="helv")

    # Page 2: Work Study & Satisfactory Academic Progress
    p2 = doc.new_page(width=595, height=842)
    text_p2 = (
        "SAMPLE UNIVERSITY — FINANCIAL AID & STUDENT ACCOUNTS (PAGE 2)\n\n"
        "SECTION: Federal Work-Study Program Guidelines\n"
        "Federal Work-Study (FWS) provides part-time employment opportunities for undergraduate and graduate "
        "students with financial need. Students awarded Work-Study may work up to 20 hours per week while "
        "classes are in session and up to 40 hours during recognized university breaks. Starting pay for campus "
        "work-study positions is $15.50 per hour. FWS earnings are paid bi-weekly directly to the student via "
        "direct deposit and are not deducted automatically from the student's tuition bill.\n\n"
        "SECTION: Satisfactory Academic Progress (SAP) Requirements\n"
        "To maintain eligibility for federal, state, and institutional financial aid, students must satisfy "
        "three Satisfactory Academic Progress standards evaluated at the end of each academic year:\n"
        "1. Minimum Cumulative GPA: Undergraduates must maintain a minimum 2.0 cumulative GPA.\n"
        "2. Pace of Completion: Students must successfully complete at least 67% of all cumulative attempted credits.\n"
        "3. Maximum Timeframe: Students must complete their degree within 150% of published program length "
        "(typically 180 attempted credit hours for a standard 120-credit bachelor's degree).\n"
        "Students who fail to meet SAP standards are placed on Financial Aid Suspension but may submit an appeal "
        "with documented extenuating circumstances before July 15."
    )
    p2.insert_textbox(rect, text_p2, fontsize=11, fontname="helv")

    pdf_path = SAMPLE_DIR / "student_financial_services.pdf"
    doc.save(str(pdf_path))
    doc.close()
    print(f"Generated PDF: {pdf_path}")

if __name__ == "__main__":
    generate_docx()
    generate_pdf()
